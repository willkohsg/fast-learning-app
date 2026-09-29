# -*- coding: utf-8 -*-
"""Fuhua-parallel SET 2 - Sec 3 G3 EOY A-Math paper, 80 marks, 2 hours.

Equation of circles is not examined in this syllabus, so the 12-mark spine is
plain coordinate geometry (perpendicular bisector, isosceles proof, area).

Every piece of mathematics is native OMML (see omml.py).
"""
import os
import sys

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from omml import add_mixed, omath                                # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
IMGDIR = os.path.join(HERE, "assets")
RIGHT_TAB = Cm(16.0)

doc = Document()
for s in doc.sections:
    s.page_width = Cm(21.0); s.page_height = Cm(29.7)       # A4
    s.top_margin = Cm(2); s.bottom_margin = Cm(2)
    s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)
st = doc.styles['Normal']
st.font.name = 'Times New Roman'; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.paragraph_format.space_after = Pt(0)


def p(text="", bold=False, italic=False, align=None, size=12, indent=None):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    if align == 'c':
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if indent is not None:
        par.paragraph_format.left_indent = Cm(indent)
    add_mixed(par, text)
    for r in par.runs:
        r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    return par


def gap(n=1):
    for _ in range(n):
        doc.add_paragraph().paragraph_format.space_after = Pt(0)


def img(name, width_cm):
    par = doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(0)
    par.add_run().add_picture(os.path.join(IMGDIR, name), width=Cm(width_cm))
    return par


def _marked(par, text, marks):
    add_mixed(par, text)
    if marks:
        par.add_run("\t[%s]" % marks)


def q(num, text, marks, indent=1.5):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.left_indent = Cm(indent)
    par.paragraph_format.first_line_indent = Cm(-indent)
    tabs = par.paragraph_format.tab_stops
    tabs.add_tab_stop(Cm(indent)); tabs.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
    par.add_run(("%s\t" % num) if num else "\t")
    _marked(par, text, marks)
    return par


def part(label, text, marks, indent=2.6):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.left_indent = Cm(indent)
    par.paragraph_format.first_line_indent = Cm(-1.1)
    tabs = par.paragraph_format.tab_stops
    tabs.add_tab_stop(Cm(indent)); tabs.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
    par.add_run("%s\t" % label).bold = True
    _marked(par, text, marks)
    return par


def eq(latex):
    par = doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(0)
    par._p.append(omath(latex))
    return par


def pagebreak():
    doc.add_page_break()


# ------------------------------------------------------------------ cover
p("SECONDARY 3 G3", bold=True, align='c', size=15)
p("END-OF-YEAR EXAMINATION 2026", bold=True, align='c', size=15)
gap(1)
p("ADDITIONAL MATHEMATICS", bold=True, align='c', size=14)
p("4049", bold=True, align='c', size=14)
par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(0)
par.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
par.add_run("Paper 1\t2 hours").bold = True
gap(2)
p("CANDIDATE NAME:  ______________________________________________")
gap(1)
p("CLASS:  3 ________________            INDEX NUMBER:  ____________")
gap(1)
p("READ THESE INSTRUCTIONS FIRST", bold=True)
p("Write your full name, class and index number in the spaces above.")
p("Write in dark blue or black pen in the space provided for each question.")
p("You may use an HB pencil for any diagrams or graphs.")
p("Do not use staples, paper clips, highlighters, glue or correction fluid.")
gap(1)
p("Answer all the questions.")
p("Give non-exact numerical answers correct to 3 significant figures, or 1 decimal place "
  "in the case of angles in degrees, unless a different level of accuracy is specified "
  "in the question.")
p("The use of an approved scientific calculator is expected, where appropriate.")
p("You are reminded of the need for clear presentation in your answers.")
p("The number of marks is given in brackets [ ] at the end of each question or part question.")
p("The total number of marks for this paper is 80.")
gap(1)

marks_list = [4, 6, 7, 8, 7, 7, 7, 12, 6, 6, 10]
tbl = doc.add_table(rows=8, cols=6)
tbl.style = 'Table Grid'; tbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
hdr = tbl.rows[0].cells; hdr[0].merge(hdr[5])
hp = tbl.rows[0].cells[0].paragraphs[0]
hp.add_run("For Examiner's Use").bold = True
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER


def fill(cell, text, bold=False):
    pp = cell.paragraphs[0]
    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.space_after = Pt(0)
    r = pp.add_run(text); r.bold = bold; r.font.size = Pt(10)


for i in range(6):
    cells = tbl.rows[i + 1].cells
    fill(cells[0], "Q%d" % (i + 1), bold=True)
    fill(cells[1], str(marks_list[i])); fill(cells[2], "")
    j = i + 6
    if j < 11:
        fill(cells[3], "Q%d" % (j + 1), bold=True)
        fill(cells[4], str(marks_list[j]))
    else:
        fill(cells[3], "Total", bold=True); fill(cells[4], "80", bold=True)
    fill(cells[5], "")
fill(tbl.rows[7].cells[0], ""); fill(tbl.rows[7].cells[1], ""); fill(tbl.rows[7].cells[2], "")
fill(tbl.rows[7].cells[3], "Total", bold=True); fill(tbl.rows[7].cells[4], "80", bold=True)
fill(tbl.rows[7].cells[5], "")
for row in tbl.rows:
    for kk, w in enumerate((2.0, 1.8, 2.4, 2.0, 1.8, 2.4)):
        row.cells[kk].width = Cm(w)
gap(1)
p("This document consists of 16 printed pages including the cover page.", align='c', size=10)
pagebreak()

# --------------------------------------------------------- formula sheet
p("Mathematical Formulae", bold=True, italic=True, align='c')
gap(1)
p("1.  ALGEBRA", bold=True, align='c')
gap(1)
p("Quadratic Equation", italic=True)
p("For the equation  $ax^2 + bx + c = 0$,", align='c')
eq(r'x = \frac{-b \pm \sqrt{b^2 - 4ac}}{2a}')
gap(1)
p("Binomial Expansion", italic=True)
eq(r'(a + b)^n = a^n + \binom{n}{1}a^{n-1}b + \binom{n}{2}a^{n-2}b^2 '
   r'+ \ldots + \binom{n}{r}a^{n-r}b^r + \ldots + b^n')
p("where n is a positive integer and", align='c')
eq(r'\binom{n}{r} = \frac{n!}{r!(n - r)!} = \frac{n(n-1)\ldots(n-r+1)}{r!}')
gap(2)
p("2.  TRIGONOMETRY", bold=True, align='c')
gap(1)
p("Identities", italic=True)
for e in [r'\sin^2 A + \cos^2 A = 1', r'\sec^2 A = 1 + \tan^2 A',
          r'\cosec^2 A = 1 + \cot^2 A',
          r'\sin(A \pm B) = \sin A\cos B \pm \cos A\sin B',
          r'\cos(A \pm B) = \cos A\cos B \mp \sin A\sin B',
          r'\tan(A \pm B) = \frac{\tan A \pm \tan B}{1 \mp \tan A\tan B}',
          r'\sin 2A = 2\sin A\cos A',
          r'\cos 2A = \cos^2 A - \sin^2 A = 2\cos^2 A - 1 = 1 - 2\sin^2 A',
          r'\tan 2A = \frac{2\tan A}{1 - \tan^2 A}']:
    eq(e)
gap(1)
p("Formulae for ΔABC", italic=True)
for e in [r'\frac{a}{\sin A} = \frac{b}{\sin B} = \frac{c}{\sin C}',
          r'a^2 = b^2 + c^2 - 2bc\cos A', r'\Delta = \frac{1}{2}ab\sin C']:
    eq(e)
pagebreak()

# -------------------------------------------------------------- questions
p("Answer all the questions.", bold=True, align='c')
gap(1)

q("1", r"The line $y = mx - 2$ meets the curve $y = x^2 - 3x + 7$ at two distinct points. "
        r"Find the range of values of $m$.", 4)
pagebreak()

q("2", "", None)
part("(a)", r"Express $\dfrac{2\sqrt{3} + 5}{3 - \sqrt{3}}$ in the form $a + b\sqrt{3}$, "
             r"where $a$ and $b$ are rational numbers.", 3)
gap(9)
part("(b)", r"Solve the equation $\sqrt{2x + 3} = x - 1$, giving your answer in exact form.", 3)
pagebreak()

q("3", r"It is given that $f(x) = 2x^3 + ax^2 + bx + 6$, where $a$ and $b$ are constants. "
        r"$(x + 2)$ is a factor of $f(x)$, and $f(3) = 0$.", None)
gap(1)
part("(a)", r"Find the value of $a$ and of $b$, and hence factorise $f(x)$ completely.", 4)
gap(10)
part("(b)", r"Hence solve the equation $2\left(2^{3y}\right) - 3\left(2^{2y}\right) "
             r"- 11\left(2^{y}\right) + 6 = 0$, giving your answers in exact form.", 3)
pagebreak()

q("4", "", None)
part("(a)", r"Show that $x^3 - x^2 + 4x - 4 = (x - 1)\left(x^2 + 4\right)$.", 2)
gap(6)
part("(b)", r"Hence express $\dfrac{2x^3 + 2x^2 + 5x + 6}{x^3 - x^2 + 4x - 4}$ "
             r"in partial fractions.", 6)
pagebreak()

q("5", "", None)
part("(a)", r"Solve the equation $\log_{3}(x + 2) + \log_{3}(x - 4) = 3$.", 3)
gap(9)
part("(b)", r"Solve the equation $4\log_{9} x - \log_{x} 3 = 1$, giving your answers "
             r"in exact form.", 4)
pagebreak()

q("6", r"The mass, $M$ grams, of a radioactive sample $t$ days after it was first "
        r"recorded is given by $M = M_{0}e^{-kt}$, where $M_{0}$ and $k$ are constants. "
        r"The mass is $500$ g when $t = 2$, and $320$ g when $t = 10$.", None)
gap(1)
part("(a)", r"Find the value of $k$ and of $M_{0}$.", 4)
gap(10)
part("(b)", r"Find the least number of complete days, measured from $t = 0$, after which "
             r"the mass first falls below $100$ g.", 3)
pagebreak()

q("7", "", None)
part("(a)", r"In the expansion of $\left(2 + kx\right)^{6}$, where $k$ is a non-zero "
             r"constant, the coefficient of $x^3$ is twice the coefficient of $x^2$. "
             r"Find the value of $k$.", 4)
gap(10)
part("(b)", r"Hence find the coefficient of $x^2$ in the expansion of "
             r"$\left(1 - 2x\right)\left(2 + 3x\right)^{6}$.", 3)
pagebreak()

q("8", r"The points $A(-2,\ 3)$ and $B(6,\ 7)$ are joined by a straight line. "
        r"Solutions by accurate drawing will not be accepted.", None)
gap(1)
part("(a)", r"Find the equation of the perpendicular bisector of $AB$.", 3)
gap(8)
part("(b)", r"The perpendicular bisector of $AB$ cuts the $x$-axis at the point $C$. "
             r"Find the coordinates of $C$.", 2)
gap(6)
pagebreak()
part("(c)", r"Explain why triangle $ABC$ is isosceles.", 2)
gap(6)
part("(d)", r"The point $D$ lies on the perpendicular bisector of $AB$ and the area of "
             r"triangle $ABD$ is $30$ units$^{2}$. Find the possible coordinates of $D$.", 5)
pagebreak()

q("9", r"The variables $x$ and $y$ are related by the equation $y = ax^{b}$, where $a$ "
        r"and $b$ are constants. The table shows measured values of $x$ and $y$.", None)
gap(1)

t2 = doc.add_table(rows=2, cols=6)
t2.style = 'Table Grid'; t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
for i, row in enumerate([["x", "1", "4", "16", "36", "100"],
                         ["y", "3.00", "6.00", "12.0", "18.0", "30.0"]]):
    for j, v in enumerate(row):
        c = t2.rows[i].cells[j]
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
        r = c.paragraphs[0].add_run(v); r.italic = (j == 0)
gap(1)
part("(a)", r"Explain how a straight line graph may be drawn to represent this equation, "
             r"and draw it for the given data on the grid below.", 3)
gap(1)
img("graphpaper.png", 14.0)
pagebreak()
part("(b)", r"Use your graph to estimate the value of $a$ and of $b$.", 3)
pagebreak()

q("10", r"Given that $\sin C = q$, where $C$ is reflex and $180^{\circ} < C < 270^{\circ}$, "
         r"express each of the following in terms of $q$.", None)
gap(1)
part("(i)", r"$\cos C$", 2)
gap(7)
part("(ii)", r"$\tan C$", 2)
gap(7)
part("(iii)", r"$\sec\left(360^{\circ} - C\right)$", 2)
pagebreak()

q("11", "", None)
part("(a)", r"Prove the identity "
             r"$\dfrac{1 - \cos\theta}{\sin\theta} + \dfrac{\sin\theta}{1 - \cos\theta} "
             r"= 2\cosec\theta$.", 3)
gap(10)
part("(b)", r"Hence solve the equation "
             r"$\dfrac{1 - \cos\theta}{\sin\theta} + \dfrac{\sin\theta}{1 - \cos\theta} = 4$ "
             r"for $0^{\circ} \le \theta \le 360^{\circ}$.", 3)
pagebreak()
part("(c)", r"Sketch the graph of $y = 2 + 3\sin 2x$ for $0^{\circ} \le x \le 360^{\circ}$, "
             r"stating the amplitude and the period of the curve.", 4)
gap(1)
img("q11axes_deg.png", 13.0)
gap(1)

p("----- End of Paper -----", bold=True, align='c')

out = os.path.join(os.path.dirname(HERE),
                   "Fuhua_Parallel_Set2_Sec3_G3_EOY_AMath_2026_Paper.docx")
doc.save(out)
print("saved", out)
