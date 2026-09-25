# -*- coding: utf-8 -*-
"""Build the worked-solutions / marking scheme for the Sec 3 EOY paper.

House format (CLG marking scheme):
    black  - question text
    blue   - step-by-step working
    red    - M/A/B mark allocation and examiner's notes

Mathematics is emitted as real OMML equation objects via omml.py.
"""
import os
import sys

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from omml import omath, colour                                  # noqa: E402

BLACK, BLUE, RED = '000000', '1F4E9C', 'C00000'
RGB = {BLACK: RGBColor(0, 0, 0), BLUE: RGBColor(0x1F, 0x4E, 0x9C),
       RED: RGBColor(0xC0, 0, 0)}
RIGHT_TAB = Cm(16.0)

doc = Document()
for s in doc.sections:
    s.page_width = Cm(21.0)      # A4 - python-docx defaults to US Letter
    s.page_height = Cm(29.7)
    s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.8)
    s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)

st = doc.styles['Normal']
st.font.name = 'Times New Roman'; st.font.size = Pt(11.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.paragraph_format.space_after = Pt(0)

_total = 0


def _emit(par, text, hexcol, bold=False, italic=False, size=11.5):
    """Write text into par; $...$ spans become OMML, all of it in hexcol."""
    import re
    for k, chunk in enumerate(re.split(r'(?<!\\)\$', text)):
        if not chunk:
            continue
        if k % 2:
            par._p.append(colour(omath(chunk), hexcol))
        else:
            r = par.add_run(chunk.replace('\\$', '$'))
            r.font.color.rgb = RGB[hexcol]
            r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    return par


def _para(indent=0.0, space_before=0, space_after=0):
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Cm(indent)
    par.paragraph_format.space_before = Pt(space_before)
    par.paragraph_format.space_after = Pt(space_after)
    par.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
    return par


def qhead(num, topic, marks):
    par = _para(space_before=14, space_after=3)
    r = par.add_run("Question %s" % num)
    r.bold = True; r.font.size = Pt(13); r.font.color.rgb = RGB[BLACK]
    r = par.add_run("   —   %s" % topic)
    r.bold = False; r.italic = True; r.font.size = Pt(11); r.font.color.rgb = RGB[BLACK]
    r = par.add_run("\t[%d]" % marks)
    r.bold = True; r.font.size = Pt(12); r.font.color.rgb = RGB[BLACK]
    # rule under the heading
    pbdr = par._p.get_or_add_pPr()
    from docx.oxml import OxmlElement
    bdr = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), '6')
    bot.set(qn('w:space'), '2'); bot.set(qn('w:color'), '000000')
    bdr.append(bot); pbdr.append(bdr)


def qtext(text, indent=0.0):
    """Question wording - black."""
    _emit(_para(indent, space_after=4), text, BLACK)


def part(label):
    par = _para(0.3, space_before=6, space_after=2)
    _emit(par, label, BLACK, bold=True)


def work(text, mark=None, indent=0.9):
    """A working step - blue - with an optional red mark code at the margin."""
    global _total
    par = _para(indent, space_after=2)
    _emit(par, text, BLUE)
    if mark:
        r = par.add_run("\t%s" % mark)
        r.bold = True; r.font.color.rgb = RGB[RED]; r.font.size = Pt(11)
        for code in mark.split():
            if code[:1] in "MAB" and code[1:].isdigit():
                _total += int(code[1:])


def note(text, indent=0.9):
    """Examiner's note - red."""
    par = _para(indent, space_after=3)
    r = par.add_run("Note:  ")
    r.bold = True; r.italic = True; r.font.size = Pt(10.5); r.font.color.rgb = RGB[RED]
    _emit(par, text, RED, italic=True, size=10.5)


def ans(text, mark=None, indent=0.9):
    global _total
    par = _para(indent, space_before=2, space_after=2)
    r = par.add_run("Answer:  ")
    r.bold = True; r.font.color.rgb = RGB[BLUE]
    _emit(par, text, BLUE, bold=True)
    if mark:
        r = par.add_run("\t%s" % mark)
        r.bold = True; r.font.color.rgb = RGB[RED]; r.font.size = Pt(11)
        for code in mark.split():
            if code[:1] in "MAB" and code[1:].isdigit():
                _total += int(code[1:])


# ----------------------------------------------------------------- cover
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("SECONDARY 3 EXPRESS  —  END-OF-YEAR EXAMINATION")
r.bold = True; r.font.size = Pt(15)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ADDITIONAL MATHEMATICS 4049"); r.bold = True; r.font.size = Pt(13)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("WORKED SOLUTIONS & MARKING SCHEME"); r.bold = True; r.font.size = Pt(14)
doc.add_paragraph()

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for label, col in (("Question", BLACK), ("  •  ", BLACK),
                   ("Working", BLUE), ("  •  ", BLACK),
                   ("Marks & examiner's notes", RED)):
    r = p.add_run(label); r.bold = True; r.font.color.rgb = RGB[col]
doc.add_paragraph()

p = doc.add_paragraph()
r = p.add_run("Mark codes:   "); r.bold = True
for code, meaning in (("M", "method mark — awarded for a correct method, even if the "
                            "arithmetic that follows is wrong"),
                      ("A", "accuracy mark — awarded for a correct answer or correct "
                            "intermediate value; dependent on the preceding M mark"),
                      ("B", "independent mark — awarded for a correct statement, value or "
                            "result on its own, with no method required")):
    q = doc.add_paragraph()
    q.paragraph_format.left_indent = Cm(1.2)
    q.paragraph_format.first_line_indent = Cm(-0.8)
    q.paragraph_format.space_after = Pt(2)
    r = q.add_run("%s   " % code); r.bold = True; r.font.color.rgb = RGB[RED]
    r = q.add_run(meaning); r.font.size = Pt(11)
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6)
r = p.add_run("Total: 90 marks.  Follow-through (ft) applies throughout: a correct method "
              "applied to a candidate's own earlier incorrect value still earns the M mark.")
r.italic = True; r.font.size = Pt(11)

doc.add_page_break()

# ================================================================== Q1
qhead(1, "Surds — area of a triangle", 3)
qtext(r"The triangle $DEF$ has an area of $\frac{1}{4}(7\sqrt{3} + 9)$ cm². The length of "
      r"$DE$ is $(\sqrt{3} + 2)$ cm and angle $EDF$ is $60\circ$. Find, without using a "
      r"calculator, the length of $DF$, in cm, in the form $(a + b\sqrt{3})$.")
work(r"Area $= \frac{1}{2}(DE)(DF)\sin\angle EDF$, so "
     r"$\frac{1}{2}(\sqrt{3}+2)(DF)\left(\frac{\sqrt{3}}{2}\right) "
     r"= \frac{1}{4}(7\sqrt{3}+9)$", "M1")
work(r"$\sqrt{3}(\sqrt{3}+2)(DF) = 7\sqrt{3}+9 \quad\Rightarrow\quad "
     r"(3+2\sqrt{3})(DF) = 7\sqrt{3}+9$")
work(r"$DF = \frac{7\sqrt{3}+9}{3+2\sqrt{3}} \times \frac{3-2\sqrt{3}}{3-2\sqrt{3}} "
     r"= \frac{3\sqrt{3}-15}{-3}$", "M1")
ans(r"$DF = 5 - \sqrt{3}$ cm, so $a = 5$ and $b = -1$.", "A1")
note(r"The $\frac{1}{2}$ and the $\sin 60\circ = \frac{\sqrt{3}}{2}$ must both appear for the "
     r"first M1. Denominator $(3)^2-(2\sqrt{3})^2 = 9-12 = -3$ is negative — candidates who "
     r"drop the sign get $\sqrt{3}-5$ and lose the A1.")

# ================================================================== Q2
qhead(2, "Surds — rationalising the denominator", 4)
part("(a)  Simplify $\\frac{7}{3-\\sqrt{2}}$ in the form $a+b\\sqrt{2}$.  [2]")
work(r"$\frac{7}{3-\sqrt{2}} \times \frac{3+\sqrt{2}}{3+\sqrt{2}} "
     r"= \frac{7(3+\sqrt{2})}{9-2} = \frac{7(3+\sqrt{2})}{7}$", "M1")
ans(r"$3 + \sqrt{2}$, so $a = 3$, $b = 1$.", "A1")
part("(b)  Given $(2+\\sqrt{3})x = 5+4\\sqrt{3}$, find $x$ in the form $a+b\\sqrt{3}$.  [2]")
work(r"$x = \frac{5+4\sqrt{3}}{2+\sqrt{3}} \times \frac{2-\sqrt{3}}{2-\sqrt{3}}$, "
     r"denominator $= 4-3 = 1$", "M1")
work(r"Numerator $= 10 - 5\sqrt{3} + 8\sqrt{3} - 12 = -2 + 3\sqrt{3}$")
ans(r"$x = -2 + 3\sqrt{3}$, so $a = -2$, $b = 3$.", "A1")
note(r"In (b) the denominator rationalises to exactly 1 — a deliberate check that "
     r"candidates still multiply out rather than guess.")

doc.add_page_break()

# ================================================================== Q3
qhead(3, "Compound angles", 4)
qtext(r"$\sin A = -\frac{1}{\sqrt{5}}$ and $\tan B = \frac{3}{4}$, where $A$ and $B$ are in "
      r"the same quadrant. Find $p$ if $\cos(A-B) = p\sqrt{5}$.")
work(r"$\sin A < 0 \Rightarrow A$ in quadrant 3 or 4;  $\tan B > 0 \Rightarrow B$ in "
     r"quadrant 1 or 3. Same quadrant $\Rightarrow$ both in quadrant 3.", "B1")
work(r"Quadrant 3: $\cos A = -\frac{2}{\sqrt{5}}$  (from $\cos^2 A = 1 - \frac{1}{5}$)")
work(r"$\tan B = \frac{3}{4}$ in quadrant 3 gives $\sin B = -\frac{3}{5}$, "
     r"$\cos B = -\frac{4}{5}$  (3–4–5 triangle, both negative)", "B1")
work(r"$\cos(A-B) = \cos A\cos B + \sin A\sin B "
     r"= \left(-\frac{2}{\sqrt{5}}\right)\left(-\frac{4}{5}\right) "
     r"+ \left(-\frac{1}{\sqrt{5}}\right)\left(-\frac{3}{5}\right)$", "M1")
work(r"$= \frac{8}{5\sqrt{5}} + \frac{3}{5\sqrt{5}} = \frac{11}{5\sqrt{5}} "
     r"= \frac{11\sqrt{5}}{25}$")
ans(r"$p = \frac{11}{25}$", "A1")
note(r"Identifying quadrant 3 is the whole question — both signs negative. A candidate who "
     r"takes quadrant 1 gets $p = \frac{11}{25}$ by luck of sign pairing but cannot earn the "
     r"two B marks.")

# ================================================================== Q4
qhead(4, "Completing the square and the discriminant", 4)
part("(i)  Express $y = -2x^2 - 12x + 5$ as $y = -a(x+b)^2 + c$.  [2]")
work(r"$-2x^2-12x+5 = -2(x^2+6x) + 5 = -2\left[(x+3)^2 - 9\right] + 5$", "M1")
ans(r"$y = -2(x+3)^2 + 23$, so $a=2$, $b=3$, $c=23$.", "A1")
part("(ii)  Range of $k$ for which $-2x^2-12x+5=k$ has no two distinct roots.  [2]")
work(r"$-2(x+3)^2 + 23 = k \quad\Rightarrow\quad (x+3)^2 = \frac{23-k}{2}$")
work(r"Two distinct roots need $\frac{23-k}{2} > 0$; so NO two distinct roots when "
     r"$\frac{23-k}{2} \le 0$", "M1")
ans(r"$k \ge 23$", "A1")
note(r"Equivalent by discriminant: $2x^2+12x+(k-5)=0$, $b^2-4ac = 184-8k \le 0$. Accept "
     r"either route. $k = 23$ MUST be included (equal roots are not two distinct "
     r"roots) — a strict inequality loses the A1.")

doc.add_page_break()

# ================================================================== Q5
qhead(5, "Discriminant — line meets curve", 5)
qtext(r"Show that $y = x-1$ intersects $y = nx^2+3x-n$ at two distinct points for all real $n$.")
work(r"Equate: $x - 1 = nx^2 + 3x - n$", "M1")
work(r"$nx^2 + 2x + (1-n) = 0$", "A1")
work(r"Discriminant $= 2^2 - 4(n)(1-n) = 4 - 4n + 4n^2 = 4(n^2 - n + 1)$", "M1")
work(r"$n^2 - n + 1 = \left(n - \frac{1}{2}\right)^2 + \frac{3}{4}$", "M1")
work(r"$\left(n-\frac{1}{2}\right)^2 \ge 0$, so $n^2-n+1 \ge \frac{3}{4} > 0$, hence "
     r"discriminant $> 0$ for all real $n$ — two distinct points. Shown.", "A1")
note(r"The final A1 needs the explicit statement that the discriminant is positive — not "
     r"merely that it is non-zero. Completing the square (or showing the discriminant of "
     r"$n^2-n+1$ is $-3 < 0$) is required; asserting positivity without justification "
     r"scores M1 only.")
note(r"Strictly the result needs $n \ne 0$, since $n = 0$ makes the curve a straight line "
     r"meeting $y = x-1$ once. The question's wording (“the curve”) presumes "
     r"$n \ne 0$; do not penalise candidates either way.")

# ================================================================== Q6
qhead(6, "Logarithms — change of base", 5)
qtext(r"Solve $\log_2(x-2) + 2\log_4(x-3) = \frac{1}{2}\log_3 9$.")
work(r"RHS: $\log_3 9 = 2$, so $\frac{1}{2}\log_3 9 = 1$", "B1")
work(r"$2\log_4(x-3) = 2 \cdot \frac{\log_2(x-3)}{\log_2 4} "
     r"= 2 \cdot \frac{\log_2(x-3)}{2} = \log_2(x-3)$", "M1")
work(r"$\log_2(x-2) + \log_2(x-3) = 1 \quad\Rightarrow\quad \log_2\left[(x-2)(x-3)\right] = 1$")
work(r"$(x-2)(x-3) = 2 \quad\Rightarrow\quad x^2 - 5x + 4 = 0 "
     r"\quad\Rightarrow\quad (x-1)(x-4) = 0$", "M1")
work(r"$x = 1$ or $x = 4$", "A1")
work(r"Domain: need $x-2>0$ AND $x-3>0$, i.e. $x > 3$. Reject $x = 1$.", "A1")
ans(r"$x = 4$")
note(r"The final A1 is for the rejection — it must be justified by the domain, not simply "
     r"omitted. A candidate giving both roots scores a maximum of 4 out of 5.")

doc.add_page_break()

# ================================================================== Q7
qhead(7, "Partial fractions", 6)
part("(i)  Explain why the form $\\frac{A}{x+1} + \\frac{B}{x^2+4}$ cannot be used.  [2]")
work(r"$x^2+4$ is an irreducible quadratic factor (it has no real linear factors, since "
     r"$x^2 = -4$ has no real solution).", "B1")
work(r"A factor of degree 2 therefore requires a numerator of degree 1, i.e. $Bx+C$, not a "
     r"constant. With only $A$ and $B$ there are two unknowns but three coefficients "
     r"($x^2$, $x$, constant) to match, so no such $A, B$ exist.", "B1")
part("(ii)  Express $\\frac{3x^2+3x+10}{(x+1)(x^2+4)}$ in partial fractions.  [4]")
work(r"Let $\frac{3x^2+3x+10}{(x+1)(x^2+4)} = \frac{A}{x+1} + \frac{Bx+C}{x^2+4}$", "M1")
work(r"$3x^2+3x+10 = A(x^2+4) + (Bx+C)(x+1)$")
work(r"$x = -1$:  $3-3+10 = A(1+4) \quad\Rightarrow\quad 10 = 5A "
     r"\quad\Rightarrow\quad A = 2$", "M1")
work(r"Compare $x^2$:  $3 = A + B \quad\Rightarrow\quad B = 1$", "A1")
work(r"Compare constants:  $10 = 4A + C = 8 + C \quad\Rightarrow\quad C = 2$", "A1")
ans(r"$\frac{2}{x+1} + \frac{x+2}{x^2+4}$")
note(r"Check by recombining: $2(x^2+4) + (x+2)(x+1) = 2x^2+8+x^2+3x+2 = 3x^2+3x+10$. "
     r"✓  Candidates who write $\frac{B}{x^2+4}$ despite part (i) can score at most the "
     r"first M1 in (ii).")

doc.add_page_break()

# ================================================================== Q8
qhead(8, "Coordinate geometry — parallelogram", 7)
qtext(r"$A(2,3)$, $B(-1,-1)$, $C(3,1)$.")
part("(i)  Show $AC \\perp BC$.  [3]")
work(r"$m_{AC} = \frac{1-3}{3-2} = -2$", "M1")
work(r"$m_{BC} = \frac{1-(-1)}{3-(-1)} = \frac{2}{4} = \frac{1}{2}$", "M1")
work(r"$m_{AC} \times m_{BC} = -2 \times \frac{1}{2} = -1$, therefore $AC \perp BC$. Shown.",
     "A1")
part("(ii)  $ABCD$ is a parallelogram; find $D$.  [2]")
work(r"Diagonals of a parallelogram bisect each other, so midpoint $AC$ = midpoint $BD$.", "M1")
work(r"Midpoint $AC = \left(\frac{2+3}{2}, \frac{3+1}{2}\right) "
     r"= \left(\frac{5}{2}, 2\right)$")
work(r"$\frac{-1+x_D}{2} = \frac{5}{2} \Rightarrow x_D = 6$;  "
     r"$\frac{-1+y_D}{2} = 2 \Rightarrow y_D = 5$")
ans(r"$D(6,\,5)$", "A1")
part("(iii)  Find the area of $ABCD$.  [2]")
work(r"Right angle at $C$, so area $= 2 \times$ area of $\triangle ABC "
     r"= 2 \times \frac{1}{2}(AC)(BC) = (AC)(BC)$", "M1")
work(r"$AC = \sqrt{1^2+(-2)^2} = \sqrt{5}$,  $BC = \sqrt{4^2+2^2} = 2\sqrt{5}$")
ans(r"Area $= \sqrt{5} \times 2\sqrt{5} = 10$ units²", "A1")
note(r"The shoelace formula on $A(2,3), B(-1,-1), C(3,1), D(6,5)$ gives 10 as well — "
     r"accept either method. Using the perpendicularity from (i) is the intended, shorter route.")

doc.add_page_break()

# ================================================================== Q9
qhead(9, "Trigonometric graphs", 7)
qtext(r"$y = a\cos 2x + b$, $a>0$, has maximum 4 and minimum $-10$ for $0 \le x \le 2\pi$.")
part("(i)  Find $a$ and $b$.  [3]")
work(r"Maximum: $a + b = 4$;  minimum: $-a + b = -10$", "M1")
work(r"Adding: $2b = -6 \Rightarrow b = -3$", "A1")
work(r"Then $a = 4 - b = 7$", "A1")
part("(ii)  State the period and the amplitude.  [2]")
work(r"Period $= \frac{2\pi}{2} = \pi$", "B1")
work(r"Amplitude $= 7$", "B1")
note(r"Amplitude is $a$, \text{not} the distance from maximum to minimum (14) and not "
     r"affected by the shift $b$. This is the most common error on this question.")
part("(iii)  Sketch $y = 7\\cos 2x - 3$ for $0 \\le x \\le 2\\pi$.  [2]")
work(r"Cosine shape, 2 complete cycles over $0 \le x \le 2\pi$", "B1")
work(r"Maximum 4 at $x = 0,\ \pi,\ 2\pi$;  minimum $-10$ at "
     r"$x = \frac{\pi}{2},\ \frac{3\pi}{2}$; oscillating about $y = -3$", "B1")
note(r"Curve must start at a maximum (cosine, positive $a$) and the two cycles must be "
     r"visibly equal in width. Crossings of $y=-3$ occur at "
     r"$x = \frac{\pi}{4}, \frac{3\pi}{4}, \frac{5\pi}{4}, \frac{7\pi}{4}$.")

doc.add_page_break()

# ================================================================== Q10
qhead(10, "Exponential growth", 7)
qtext(r"$P = 45\,000e^{kt}$, $t$ years after the beginning of 2015. At the beginning of 2016 "
      r"the vase was valued at \$50 000.")
part("(a)  Show that $k = 0.10536$ (5 s.f.).  [3]")
work(r"$t = 1$, $P = 50\,000$:  $50\,000 = 45\,000e^{k}$", "M1")
work(r"$e^{k} = \frac{50\,000}{45\,000} = \frac{10}{9}$", "M1")
work(r"$k = \ln\frac{10}{9} = 0.1053605\ldots = 0.10536$ (5 s.f.). Shown.", "A1")
part("(b)  Find the year in which the investment triples.  [4]")
work(r"Tripled: $P = 3 \times 45\,000 = 135\,000$", "M1")
work(r"$135\,000 = 45\,000e^{kt} \quad\Rightarrow\quad e^{kt} = 3$", "M1")
work(r"$t = \frac{\ln 3}{0.1053605} = \frac{1.098612}{0.1053605} = 10.427\ldots$", "M1")
work(r"$t \approx 10.43$ years after the beginning of 2015", )
ans(r"During 2025", "A1")
note(r"$t = 10$ is the beginning of 2025 and $t = 11$ the beginning of 2026, so "
     r"$t = 10.43$ falls inside 2025. The A1 is for the YEAR, not for $t$ — a "
     r"candidate who stops at 10.4 years scores M1 M1 M1 only.")
note(r"Use the exact $k = \ln\frac{10}{9}$ where possible; using the rounded 0.10536 is "
     r"accepted here because the question supplied it.")

doc.add_page_break()

# ================================================================== Q11
qhead(11, "Binomial theorem", 8)
part("(a)  Term independent of $x$ in $3x^3\\left(2x-\\frac{1}{4x^2}\\right)^{12}$.  [3]")
work(r"General term of $\left(2x-\frac{1}{4x^2}\right)^{12}$: "
     r"$\binom{12}{r}(2x)^{12-r}\left(-\frac{1}{4x^2}\right)^{r}$", "M1")
work(r"$= \binom{12}{r}\,2^{12-r}\,(-1)^r\,4^{-r}\,x^{12-3r}$")
work(r"With the factor $3x^3$ the power of $x$ is $3 + 12 - 3r = 15 - 3r$; "
     r"set $15-3r = 0 \Rightarrow r = 5$", "M1")
work(r"Term $= 3\binom{12}{5}2^{7}(-1)^5 4^{-5} "
     r"= 3 \times 792 \times 128 \times \frac{-1}{1024} = -3 \times 99$")
ans(r"$-297$", "A1")
note(r"$\binom{12}{5} = 792$ and $\frac{128}{1024} = \frac{1}{8}$. The sign comes from "
     r"$(-1)^5$ — losing it gives $+297$ and costs the A1.")
part("(b)  In $(2+mx)^{15}$ the ratio of the coefficient of $x^7$ to that of $x^9$ is "
     "$\\frac{4}{7}$. Find $m$.  [5]")
work(r"Coefficient of $x^7$: $\binom{15}{7}2^{8}m^{7}$", "M1")
work(r"Coefficient of $x^9$: $\binom{15}{9}2^{6}m^{9}$", "M1")
work(r"$\frac{\binom{15}{7}2^{8}m^{7}}{\binom{15}{9}2^{6}m^{9}} = \frac{4}{7}$", "M1")
work(r"$\frac{6435}{5005} = \frac{9}{7}$, so $\frac{9}{7}\cdot\frac{4}{m^2} = \frac{4}{7} "
     r"\quad\Rightarrow\quad \frac{36}{7m^2} = \frac{4}{7}$", "M1")
work(r"$m^2 = 9$", )
ans(r"$m = 3$ or $m = -3$", "A1")
note(r"Both values are required — the ratio involves $m^2$, so no root can be rejected. "
     r"Giving only $m = 3$ scores 4 out of 5.")

doc.add_page_break()

# ================================================================== Q12
qhead(12, "Polynomials — remainder and factor theorems", 10)
part("(a)  $f(x)=x^3+x^2+ax+2b$ has factor $(x+3)$; $g(x)=x^3-4x^2-ax-b$ leaves remainder "
     "42 on division by $(x-5)$. Find $a$ and $b$.  [5]")
work(r"$(x+3)$ a factor $\Rightarrow f(-3) = 0$:  $-27 + 9 - 3a + 2b = 0$", "M1")
work(r"$2b = 3a + 18$  …(1)", "A1")
work(r"$g(5) = 42$:  $125 - 100 - 5a - b = 42$", "M1")
work(r"$b = -5a - 17$  …(2)", "A1")
work(r"Substituting (2) into (1): $2(-5a-17) = 3a+18 \Rightarrow -13a = 52$")
ans(r"$a = -4$,  $b = 3$", "A1")
note(r"Check: (1) gives $2(3) = 6$ and $3(-4)+18 = 6$ ✓. Sign slips in $f(-3)$ "
     r"($(-3)^3 = -27$, $(-3)^2 = +9$) are the usual cause of lost A marks.")
part("(b)  Solve $x^3-x^2-7x+7=0$, giving the roots in exact form.  [5]")
work(r"Group: $x^2(x-1) - 7(x-1) = 0$", "M1")
work(r"$(x-1)(x^2-7) = 0$", "A1")
work(r"$x^2 - 7 = 0 \Rightarrow x^2 = 7$", "M1")
work(r"$x = 1$", "A1")
ans(r"$x = 1$,  $x = \sqrt{7}$,  $x = -\sqrt{7}$", "A1")
note(r"“Exact form” means $\pm\sqrt{7}$ must be left in surd form — giving "
     r"$\pm 2.65$ loses the final A1. Factor theorem with trial $x=1$ is equally acceptable "
     r"in place of grouping.")

doc.add_page_break()

# ================================================================== Q13
qhead(13, "Linear law", 8)
qtext(r"$L = Ad^{\,k}$, with $d$: 9, 50, 200, 400, 900 and $L$: 8, 15, 25, 32, 43.")
part("(a)  Explain how a straight line graph can be drawn, and draw it.  [4]")
work(r"Take $\lg$ of both sides: $\lg L = \lg A + k\lg d$", "B1")
work(r"Comparing with $Y = mX + c$: plot $\lg L$ (vertical) against $\lg d$ (horizontal); "
     r"the graph is a straight line of gradient $k$ and vertical intercept $\lg A$.", "B1")
work(r"$\lg d$:  0.954,  1.699,  2.301,  2.602,  2.954", "M1")
work(r"$\lg L$:  0.903,  1.176,  1.398,  1.505,  1.633", "A1")
part("(b)  Estimate $A$ and $k$.  [4]")
work(r"Gradient $k = \frac{1.633-0.903}{2.954-0.954} = \frac{0.730}{2.000}$", "M1")
ans(r"$k \approx 0.365$", "A1")
work(r"Intercept: $\lg A = \lg L - k\lg d = 0.903 - 0.365(0.954) = 0.555$", "M1")
ans(r"$A = 10^{0.555} \approx 3.59$", "A1")
note(r"These are graphical estimates: accept $k$ in $0.35$ to $0.38$ and $A$ in $3.3$ to "
     r"$3.9$, provided the gradient is computed from two points on the DRAWN LINE "
     r"and not from two raw data points.")
note(r"Check: $3.59 \times 200^{0.365} = 3.59 \times 6.92 \approx 24.8$, against the "
     r"tabulated $L = 25$ ✓")

doc.add_page_break()

# ================================================================== Q14
IDENT = r'\frac{\cos\theta}{1-\sin\theta} - \frac{1-\sin\theta}{\cos\theta}'
qhead(14, "Trigonometric identities", 12)
part("(a)  Prove $" + IDENT + " = 2\\tan\\theta$.  [5]")
work(r"LHS $= \frac{\cos^2\theta - (1-\sin\theta)^2}{(1-\sin\theta)\cos\theta}$", "M1")
work(r"$(1-\sin\theta)^2 = 1 - 2\sin\theta + \sin^2\theta$", "M1")
work(r"Numerator $= \cos^2\theta - 1 + 2\sin\theta - \sin^2\theta$; "
     r"using $\cos^2\theta = 1-\sin^2\theta$:", "M1")
work(r"$= 1 - \sin^2\theta - 1 + 2\sin\theta - \sin^2\theta "
     r"= 2\sin\theta - 2\sin^2\theta = 2\sin\theta(1-\sin\theta)$", "M1")
work(r"LHS $= \frac{2\sin\theta(1-\sin\theta)}{(1-\sin\theta)\cos\theta} "
     r"= \frac{2\sin\theta}{\cos\theta} = 2\tan\theta = \text{RHS}$. Proved.", "A1")
note(r"Work on one side only. Candidates who cross-multiply from LHS = RHS have assumed "
     r"the result and score a maximum of 3 (the M marks for the algebra).")
part("(b)  Hence solve $" + IDENT + " = \\cot\\theta$ for $0 \\le \\theta \\le 2\\pi$.  [4]")
work(r"By (a): $2\tan\theta = \cot\theta = \frac{1}{\tan\theta}$", "M1")
work(r"$2\tan^2\theta = 1 \quad\Rightarrow\quad \tan\theta = \pm\frac{1}{\sqrt{2}}$", "M1")
work(r"Basic angle $\alpha = \tan^{-1}\frac{1}{\sqrt{2}} = 0.6155$ rad")
work(r"$\tan\theta = +\frac{1}{\sqrt{2}}$: $\theta = 0.6155$, $\pi + 0.6155 = 3.757$", "A1")
work(r"$\tan\theta = -\frac{1}{\sqrt{2}}$: $\theta = \pi - 0.6155 = 2.526$, "
     r"$2\pi - 0.6155 = 5.668$")
ans(r"$\theta = 0.615,\ 2.53,\ 3.76,\ 5.67$ rad (3 s.f.)", "A1")
note(r"All four solutions are required for the final A1. The range is in radians — "
     r"answers in degrees score a maximum of 3. Note $\theta \ne \frac{\pi}{2}, "
     r"\frac{3\pi}{2}$ (undefined) but none of the solutions violate this.")
part("(c)  Find the set of values of $k$ for which $" + IDENT + " = k(1+\\tan^2\\theta)$ "
     "has no solutions.  [3]")
work(r"By (a) the equation becomes $2\tan\theta = k(1+\tan^2\theta)$. "
     r"Let $t = \tan\theta$:", "M1")
work(r"$kt^2 - 2t + k = 0$")
work(r"No real $t$ requires discriminant $< 0$: $(-2)^2 - 4(k)(k) < 0$", "M1")
work(r"$4 - 4k^2 < 0 \quad\Rightarrow\quad k^2 > 1$")
ans(r"$k < -1$  or  $k > 1$   (equivalently $k^2 > 1$)", "A1")
note(r"$k = 0$ gives $-2t = 0$, i.e. $t = 0$, which DOES have solutions "
     r"($\theta = 0, \pi, 2\pi$) — consistent, since $0^2 < 1$. Candidates who divide "
     r"by $k$ without considering $k = 0$ should not be penalised provided the final set "
     r"is correct.")

out = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "Sec3_EOY_AMath_Solutions.docx")
doc.save(out)
print("saved", out)
print("marks accounted for:", _total)
assert _total == 90, "mark total is %d, expected 90" % _total
