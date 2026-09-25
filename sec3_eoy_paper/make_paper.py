# -*- coding: utf-8 -*-
"""Build the Sec 3 EOY Additional Mathematics paper.

All mathematics is emitted as real OMML equation objects (see omml.py); text
passed to p()/q()/part() may contain $...$ spans, which become equations.
"""
import os
import sys

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from omml import add_mixed, omath                      # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
IMGDIR = os.path.join(HERE, "assets")
RIGHT_TAB = Cm(16.0)

doc = Document()
for s in doc.sections:
    s.page_width = Cm(21.0)      # A4 - python-docx defaults to US Letter
    s.page_height = Cm(29.7)
    s.top_margin = Cm(2); s.bottom_margin = Cm(2)
    s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)

st = doc.styles['Normal']
st.font.name = 'Times New Roman'; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.paragraph_format.space_after = Pt(0)


def p(text="", bold=False, italic=False, align=None, size=12, indent=None):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    if align == 'c':
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if indent is not None:
        par.paragraph_format.left_indent = Cm(indent)
    add_mixed(par, text, {'bold': bold, 'italic': italic, 'font.size': None})
    for r in par.runs:
        r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    return par


def gap(n=1):
    for _ in range(n):
        doc.add_paragraph().paragraph_format.space_after = Pt(0)


def img(name, width_cm):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(0)
    par.add_run().add_picture(os.path.join(IMGDIR, name), width=Cm(width_cm))
    return par


def _marked(par, text, marks):
    add_mixed(par, text)
    if marks:
        par.add_run("\t[%s]" % marks)


def q(num, text, marks, indent=1.5):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.left_indent = Cm(indent)
    par.paragraph_format.first_line_indent = Cm(-indent)
    tabs = par.paragraph_format.tab_stops
    tabs.add_tab_stop(Cm(indent))
    tabs.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
    par.add_run(("%s\t" % num) if num else "\t")
    _marked(par, text, marks)
    return par


def part(label, text, marks, indent=2.6):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.left_indent = Cm(indent)
    par.paragraph_format.first_line_indent = Cm(-1.1)
    tabs = par.paragraph_format.tab_stops
    tabs.add_tab_stop(Cm(indent))
    tabs.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
    par.add_run("%s\t" % label).bold = True
    _marked(par, text, marks)
    return par


def eq(latex, indent=2.6):
    """A display equation on its own line."""
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par._p.append(omath(latex))
    return par


def pagebreak():
    doc.add_page_break()


# ------------------------------------------------------------------ cover
p("SECONDARY 3 EXPRESS", bold=True, align='c', size=15)
p("END-OF-YEAR EXAMINATION", bold=True, align='c', size=15)
gap(1)
p("ADDITIONAL MATHEMATICS", bold=True, align='c', size=14)
p("4049", bold=True, align='c', size=14)
par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(0)
par.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
par.add_run("Paper 1\t2 hours 15 minutes").bold = True
gap(2)
p("CANDIDATE NAME:  ______________________________________________")
gap(1)
p("CLASS:  3 ________________            INDEX NUMBER:  ____________")
gap(1)
p("READ THESE INSTRUCTIONS FIRST", bold=True)
p("Write your full name, class and index number in the spaces above.")
p("Write in dark blue or black pen in the space provided for each question.")
p("You may use a HB pencil for any diagrams or graphs.")
p("Do not use staples, paper clips, highlighters, glue or correction fluid.")
gap(1)
p("Answer all the questions.")
p("The number of marks is given in brackets [ ] at the end of each question or part question.")
p("If working is needed for any question, it must be shown in the space below the question.")
p("Omission of essential working will result in loss of marks.")
p("The total of the marks for this paper is 90.")
gap(1)
p("The use of an approved scientific calculator is expected, where appropriate.")
p("If the degree of accuracy is not specified in the question and if the answer is not exact, "
  "give the answer to three significant figures. Give answers in degrees to one decimal place.")
p("For π, use either your calculator value or 3.142.")
gap(1)

marks_list = [3, 4, 4, 4, 5, 5, 6, 7, 7, 7, 8, 10, 8, 12]
tbl = doc.add_table(rows=9, cols=6)
tbl.style = 'Table Grid'
tbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
hdr = tbl.rows[0].cells
hdr[0].merge(hdr[5])
hp = tbl.rows[0].cells[0].paragraphs[0]
hp.add_run("For Examiner's Use").bold = True
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER


def fill(cell, text, bold=False):
    pp = cell.paragraphs[0]
    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.space_after = Pt(0)
    r = pp.add_run(text); r.bold = bold; r.font.size = Pt(10)


for i in range(7):
    cells = tbl.rows[i + 1].cells
    fill(cells[0], "Q%d" % (i + 1), bold=True)
    fill(cells[1], str(marks_list[i]))
    fill(cells[2], "")
    j = i + 7
    fill(cells[3], "Q%d" % (j + 1), bold=True)
    fill(cells[4], str(marks_list[j]))
    fill(cells[5], "")
for k in (0, 1, 2, 5):
    fill(tbl.rows[8].cells[k], "")
fill(tbl.rows[8].cells[3], "Total", bold=True)
fill(tbl.rows[8].cells[4], "90", bold=True)
for row in tbl.rows:
    for k, w in enumerate((2.0, 1.8, 2.4, 2.0, 1.8, 2.4)):
        row.cells[k].width = Cm(w)

gap(1)
p("This document consists of 18 printed pages including the cover page.", align='c', size=10)
pagebreak()

# --------------------------------------------------------- formula sheet
p("Mathematical Formulae", bold=True, italic=True, align='c')
gap(1)
p("1.  ALGEBRA", bold=True, align='c')
gap(1)
p("Quadratic Equation", italic=True)
p("For the equation  $ax^2 + bx + c = 0$,", align='c')
eq(r'x = \frac{-b \pm \sqrt{b^2 - 4ac}}{2a}')
gap(1)
p("Binomial Expansion", italic=True)
eq(r'(a + b)^n = a^n + \binom{n}{1}a^{n-1}b + \binom{n}{2}a^{n-2}b^2 '
   r'+ \ldots + \binom{n}{r}a^{n-r}b^r + \ldots + b^n')
p("where n is a positive integer and", align='c')
eq(r'\binom{n}{r} = \frac{n!}{r!(n - r)!} = \frac{n(n-1)\ldots(n-r+1)}{r!}')
gap(2)
p("2.  TRIGONOMETRY", bold=True, align='c')
gap(1)
p("Identities", italic=True)
for e in [r'\sin^2 A + \cos^2 A = 1',
          r'\sec^2 A = 1 + \tan^2 A',
          r'\cosec^2 A = 1 + \cot^2 A',
          r'\sin(A \pm B) = \sin A\cos B \pm \cos A\sin B',
          r'\cos(A \pm B) = \cos A\cos B \mp \sin A\sin B',
          r'\tan(A \pm B) = \frac{\tan A \pm \tan B}{1 \mp \tan A\tan B}',
          r'\sin 2A = 2\sin A\cos A',
          r'\cos 2A = \cos^2 A - \sin^2 A = 2\cos^2 A - 1 = 1 - 2\sin^2 A',
          r'\tan 2A = \frac{2\tan A}{1 - \tan^2 A}']:
    eq(e)
gap(1)
p("Formulae for ΔABC", italic=True)
for e in [r'\frac{a}{\sin A} = \frac{b}{\sin B} = \frac{c}{\sin C}',
          r'a^2 = b^2 + c^2 - 2bc\cos A',
          r'\Delta = \frac{1}{2}ab\sin C']:
    eq(e)
pagebreak()

# -------------------------------------------------------------- questions
p("Answer all the questions.", bold=True, align='c')
gap(1)

q("1", r"The triangle $DEF$ has an area of $\frac{1}{4}(7\sqrt{3} + 9)$ cm². "
        r"The length of $DE$ is $(\sqrt{3} + 2)$ cm and angle $EDF$ is $60\circ$. "
        r"Find, without using a calculator, the length of $DF$, in cm, in the form "
        r"$(a + b\sqrt{3})$ where $a$ and $b$ are integers.", 3)
gap(11)

q("2", "Do not use a calculator in this question.", None)
part("(a)", r"Simplify $\frac{7}{3 - \sqrt{2}}$, giving your answer in the form "
             r"$a + b\sqrt{2}$, where $a$ and $b$ are integers.", 2)
gap(8)
part("(b)", r"Given that $(2 + \sqrt{3})x = 5 + 4\sqrt{3}$, find the value of $x$ in the form "
             r"$a + b\sqrt{3}$, where $a$ and $b$ are integers.", 2)
pagebreak()

q("3", r"It is given that $\sin A = -\frac{1}{\sqrt{5}}$ and $\tan B = \frac{3}{4}$, "
        r"where $A$ and $B$ are in the same quadrant. Find the value of $p$ if "
        r"$\cos(A - B) = p\sqrt{5}$.", 4)
pagebreak()

q("4", "", None)
part("(i)", r"Express $y = -2x^2 - 12x + 5$ in the form $y = -a(x + b)^2 + c$, "
             r"where $a$, $b$ and $c$ are constants.", 2)
gap(10)
part("(ii)", r"Hence find the range of values of the constant $k$ for which the equation "
              r"$-2x^2 - 12x + 5 = k$ does not have two distinct roots.", 2)
pagebreak()

q("5", r"Show that the line $y = x - 1$ intersects the curve $y = nx^2 + 3x - n$ "
        r"at two distinct points for all real values of $n$.", 5)
pagebreak()

q("6", r"Solve the equation $\log_2(x - 2) + 2\log_4(x - 3) = \frac{1}{2}\log_3 9$.", 5)
pagebreak()

PF = r'\frac{3x^2 + 3x + 10}{(x + 1)(x^2 + 4)}'
q("7", "", None)
part("(i)", r"Explain why $" + PF + r"$ cannot be expressed in the form", 2)
eq(r'\frac{A}{x + 1} + \frac{B}{x^2 + 4}\quad\text{, where } A \text{ and } B \text{ are constants.}')
gap(9)
part("(ii)", r"Express $" + PF + r"$ in partial fractions.", 4)
pagebreak()

q("8", r"The coordinates of the points $A$, $B$ and $C$ are $(2, 3)$, $(-1, -1)$ and "
        r"$(3, 1)$ respectively.", None)
gap(1)
part("(i)", r"Show that $AC$ is perpendicular to $BC$.", 3)
gap(8)
part("(ii)", r"Given that $ABCD$ is a parallelogram, find the coordinates of the point $D$.", 2)
gap(6)
part("(iii)", r"Find the area of the parallelogram $ABCD$.", 2)
pagebreak()

q("9", r"The curve $y = a\cos 2x + b$, where $a > 0$, has a maximum value of 4 and a "
        r"minimum value of $-10$, for $0 \le x \le 2\pi$.", None)
gap(1)
part("(i)", r"Find the value of $a$ and of $b$.", 3)
gap(7)
part("(ii)", r"State the period and the amplitude of the curve $y = a\cos 2x + b$.", 2)
gap(5)
part("(iii)", r"Sketch the curve of $y = a\cos 2x + b$ for $0 \le x \le 2\pi$ "
               r"on the axes below.", 2)
gap(1)
img("trigaxes.png", 13.5)
pagebreak()

q("10", r"Mr Goh bought an antique vase at the beginning of 2015 as an investment. "
         r"According to the seller, the vase's value, \$$P$, increases over time, $t$ years, "
         r"after purchase. The value of the vase is given by $P = 45\,000e^{kt}$, where $k$ is "
         r"a constant. An evaluation at the beginning of 2016 valued the vase at \$50 000.", None)
gap(1)
part("(a)", r"Show that $k = 0.10536$, correct to 5 significant figures.", 3)
gap(8)
part("(b)", r"Assuming that the value of the vase continues to increase at the same rate, "
             r"calculate the year in which Mr Goh's investment becomes tripled.", 4)
pagebreak()

q("11", "", None)
part("(a)", r"Find the term independent of $x$ in the expansion of "
             r"$3x^3\left(2x - \frac{1}{4x^2}\right)^{12}$.", 3)
gap(11)
part("(b)", r"In the expansion of $(2 + mx)^{15}$, where $m$ is a constant, the ratio of the "
             r"coefficient of $x^7$ to the coefficient of $x^9$ is $\frac{4}{7}$. "
             r"Find the possible values of $m$.", 5)
pagebreak()

q("12", "", None)
part("(a)", r"It is given that $f(x) = x^3 + x^2 + ax + 2b$ and "
             r"$g(x) = x^3 - 4x^2 - ax - b$, where $a$ and $b$ are constants.", None)
p(r"$(x + 3)$ is a factor of $f(x)$.", indent=3.7)
p(r"$g(x)$ leaves a remainder of 42 when divided by $(x - 5)$.", indent=3.7)
par = doc.add_paragraph()
par.paragraph_format.left_indent = Cm(3.7); par.paragraph_format.space_after = Pt(0)
par.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
add_mixed(par, r"Find the values of $a$ and of $b$.")
par.add_run("\t[5]")
gap(11)
part("(b)", r"Solve the equation $x^3 - x^2 - 7x + 7 = 0$, expressing the roots in "
             r"exact form.", 5)
pagebreak()

q("13", r"A marine monitoring buoy transmits data via Wi-Fi to a receiving station on shore. "
         r"As the signal travels over the sea, it experiences a loss in strength. The signal "
         r"loss, $L$ (in dB), depends on the distance, $d$ m, from the buoy and can be "
         r"modelled by the formula $L = Ad^k$, where $A$ and $k$ are constants.", None)
gap(1)
p("The table below shows corresponding values of $d$ and $L$.", indent=1.5)
gap(1)
t2 = doc.add_table(rows=2, cols=6)
t2.style = 'Table Grid'; t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
for i, row in enumerate([["d", "9", "50", "200", "400", "900"],
                         ["L", "8", "15", "25", "32", "43"]]):
    for j, v in enumerate(row):
        c = t2.rows[i].cells[j]
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
        r = c.paragraphs[0].add_run(v)
        r.italic = (j == 0)
gap(1)
part("(a)", r"Explain how a straight line graph can be drawn to represent the formula, "
             r"and draw it for the given data on the grid below.", 4)
gap(1)
img("graphpaper.png", 14.5)
pagebreak()
part("(b)", r"Use your graph to estimate the value of $A$ and of $k$.", 4)
pagebreak()

IDENT = r'\frac{\cos\theta}{1 - \sin\theta} - \frac{1 - \sin\theta}{\cos\theta}'
q("14", "", None)
part("(a)", r"Prove the identity", 5)
eq(IDENT + r' = 2\tan\theta')
gap(12)
part("(b)", r"Hence, solve the equation", 4)
eq(IDENT + r' = \cot\theta \quad \text{for } 0 \le \theta \le 2\pi')
pagebreak()
part("(c)", r"Use the result in part (a) to find the set of values of the constant $k$ "
             r"for which there are no solutions to the equation", 3)
eq(IDENT + r' = k(1 + \tan^2\theta)')
gap(20)
p("----- End of Paper -----", bold=True, align='c')

out = os.path.join(os.path.dirname(HERE), "Sec3_EOY_AMath_Paper.docx")
doc.save(out)
print("saved", out)
